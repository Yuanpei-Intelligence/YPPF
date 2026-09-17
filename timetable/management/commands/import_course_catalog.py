"""
``import_course_catalog <xlsx> --term <code> [--sheet <name>]``: load the
PKU-Course-Crawler workbook (``课表信息汇总+.xlsx``) into
``CourseCatalogEntry`` rows (``timetable/README.md`` §6.3).

The header row is matched by column names (any order, a few aliases
accepted). Rows go to the term named by their 学年学期 cell when that
resolves to an existing ``AcademicTerm``; rows whose cell is absent or
unrecognised go to ``--term``, rows naming a term that does not exist are
skipped and reported. Course codes stored as numbers are zero-padded to
eight digits and class numbers to two, matching the university's format.
"""
from __future__ import annotations

import zipfile
from typing import Any

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from timetable.catalog import normalise_catalog_row, parse_term_code, upsert_catalog_rows
from timetable.models import AcademicTerm

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    'term_text': ('学年学期', '学期', '开课学期'),
    'department': ('院系', '开课院系', '开课单位', '学院'),
    'course_code': ('课程号', '课程编号', '课程代码'),
    'name': ('课程名', '课程名称', '课程中文名', '中文名称'),
    'name_en': ('课程英文名', '课程英文名称', '英文名称', '英文名'),
    'class_no': ('班号', '班级号', '教学班号', '班级'),
    'audience': ('修读对象', '授课对象', '面向对象'),
    'category': ('课程类别', '课程类型', '类别'),
    'credits': ('参考学分', '学分'),
    'hours_per_week': ('周学时',),
    'total_hours': ('总学时',),
    'teacher': ('授课教师', '教师', '任课教师', '教师姓名'),
    'weeks_text': ('起止周', '周次', '上课周次', '起止周次'),
    'time_text': ('上课时间', '时间', '上课时间地点', '时间地点'),
    'note': ('备注', '说明'),
}
REQUIRED_COLUMNS = ('course_code', 'name')
HEADER_SCAN_ROWS = 20


def _header_text(value: Any) -> str:
    return ''.join(str(value or '').split())


def _column_map(row: tuple[Any, ...]) -> dict[str, int] | None:
    # {field: column index} of a header row, or None if it is not one.
    by_name: dict[str, int] = {}
    for index, cell in enumerate(row):
        text = _header_text(cell)
        if text:
            by_name.setdefault(text, index)
    columns: dict[str, int] = {}
    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in by_name:
                columns[field] = by_name[alias]
                break
    if all(field in columns for field in REQUIRED_COLUMNS):
        return columns
    return None


def _cell(value: Any, field: str) -> Any:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        if field == 'course_code':
            return f'{value:08d}'
        if field == 'class_no':
            return f'{value:02d}'
    return value


class Command(BaseCommand):
    help = ('Import the PKU course catalog workbook (ICUlizhi/PKU-Course-Crawler '
            '课表信息汇总+.xlsx) into CourseCatalogEntry rows of a term.')

    def add_arguments(self, parser):
        parser.add_argument('xlsx', help='path of the workbook')
        parser.add_argument(
            '--term', required=True,
            help='AcademicTerm code (e.g. 26-27-1) for rows whose 学年学期 is '
                 'absent or not recognised')
        parser.add_argument(
            '--sheet', default=None,
            help='worksheet name (default: the first worksheet)')

    def handle(self, *args, **options):
        default_term = AcademicTerm.objects.filter(code=options['term']).first()
        if default_term is None:
            raise CommandError(
                f'unknown term {options["term"]!r}; create it in the admin or '
                f'run timetable_seed_terms first')
        try:
            workbook = load_workbook(options['xlsx'], read_only=True, data_only=True)
        except (OSError, InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
            raise CommandError(f'cannot open workbook: {exc}')
        try:
            sheet = self._sheet(workbook, options['sheet'])
            columns, rows = self._read_rows(sheet)
        finally:
            workbook.close()
        self.stdout.write(f'sheet {sheet.title!r}: {len(rows)} data row(s), '
                          f'columns {sorted(columns)}')

        rows_by_term: dict[str, list[dict[str, Any]]] = {}
        terms: dict[str, AcademicTerm] = {default_term.code: default_term}
        unknown_terms: dict[str, int] = {}
        skipped_unkeyed = 0
        for row in rows:
            if normalise_catalog_row(row) is None:
                skipped_unkeyed += 1
                continue
            code = parse_term_code(row.pop('term_text', None)) or default_term.code
            if code not in terms:
                term = AcademicTerm.objects.filter(code=code).first()
                if term is None:
                    unknown_terms[code] = unknown_terms.get(code, 0) + 1
                    continue
                terms[code] = term
            rows_by_term.setdefault(code, []).append(row)

        total_created = total_updated = 0
        for code in sorted(rows_by_term):
            created, updated = upsert_catalog_rows(terms[code], rows_by_term[code])
            total_created += created
            total_updated += updated
            self.stdout.write(f'{code}: {len(rows_by_term[code])} row(s) -> '
                              f'created {created}, updated {updated}')
        summary = (f'done: created {total_created}, updated {total_updated}; '
                   f'skipped {skipped_unkeyed} row(s) without 课程号/课程名')
        if unknown_terms:
            detail = ', '.join(f'{code} ({count})'
                               for code, count in sorted(unknown_terms.items()))
            summary += (f', {sum(unknown_terms.values())} row(s) of unknown '
                        f'term(s): {detail}')
        self.stdout.write(summary)

    def _sheet(self, workbook, name: str | None):
        if name is None:
            if not workbook.sheetnames:
                raise CommandError('the workbook has no worksheet')
            return workbook[workbook.sheetnames[0]]
        if name not in workbook.sheetnames:
            raise CommandError(
                f'worksheet {name!r} not found; available: {workbook.sheetnames}')
        return workbook[name]

    def _read_rows(self, sheet) -> tuple[dict[str, int], list[dict[str, Any]]]:
        columns: dict[str, int] | None = None
        rows: list[dict[str, Any]] = []
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if columns is None:
                columns = _column_map(values)
                if columns is None and number >= HEADER_SCAN_ROWS:
                    break
                continue
            if all(value is None or str(value).strip() == '' for value in values):
                continue
            row: dict[str, Any] = {}
            for field, index in columns.items():
                value = values[index] if index < len(values) else None
                row[field] = _cell(value, field)
            rows.append(row)
        if columns is None:
            raise CommandError(
                f'no header row with 课程号 and 课程名 in the first {HEADER_SCAN_ROWS} rows')
        return columns, rows
