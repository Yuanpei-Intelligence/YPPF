"""
``import_exam_schedule <xlsx|csv> --term <code> [--sheet <name>] [--replace]``:
load the 教务部 exam table into ``CourseExam`` rows of a term
(``timetable/README.md`` §8.4).

The header row is matched by column names (any order, aliases accepted,
like ``import_course_catalog``). The exam time comes from a single
考试时间 cell or from 考试日期 + 开始时间/结束时间 columns; parsing lives in
``timetable.exams.parse_exam_time`` and rows it cannot read are skipped
and reported. Rows are upserted on ``(term, course_code, class_no,
start)``; other rows of the term are kept unless ``--replace`` deletes them
first. A 学年学期 column, when present, lets rows of another term be
skipped and reported.
"""
from __future__ import annotations

import csv
import zipfile
from pathlib import Path
from typing import Any

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from timetable.catalog import normalise_class_no, normalise_course_code, parse_term_code
from timetable.exams import parse_exam_time
from timetable.models import AcademicTerm, CourseExam

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    'term_text': ('学年学期', '学期', '开课学期'),
    'course_code': ('课程号', '课程编号', '课程代码'),
    'name': ('课程名', '课程名称', '课程中文名'),
    'class_no': ('班号', '班级号', '教学班号', '班级'),
    'teacher': ('教师', '授课教师', '任课教师', '教师姓名', '主讲教师'),
    'time_text': ('考试时间', '考试日期时间', '时间', '考试日期及时间'),
    'date_text': ('考试日期', '日期'),
    'start_text': ('开始时间', '起始时间'),
    'end_text': ('结束时间', '终止时间'),
    'room': ('考试地点', '教室', '考场', '地点', '考试教室'),
    'method': ('考试方式', '考核方式'),
    'note': ('备注', '说明'),
}
REQUIRED_COLUMNS = ('course_code', 'name')
TIME_COLUMNS = ('time_text', 'date_text')
HEADER_SCAN_ROWS = 20
_TEXT_LIMITS = {'name': 100, 'teacher': 80, 'room': 100, 'method': 32,
                'note': 200, 'course_code': 32, 'class_no': 8}


def _header_text(value: Any) -> str:
    return ''.join(str(value or '').split())


def _column_map(row: tuple[Any, ...]) -> dict[str, int] | None:
    # {field: column index} of a header row, or None if it is not one.
    by_name: dict[str, int] = {}
    for index, cell in enumerate(row):
        text = _header_text(cell)
        if text:
            by_name.setdefault(text, index)
    columns: dict[str, int] = {}
    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in by_name:
                columns[field] = by_name[alias]
                break
    if (all(field in columns for field in REQUIRED_COLUMNS)
            and any(field in columns for field in TIME_COLUMNS)):
        return columns
    return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, str)):
        return ' '.join(str(value).split())
    # datetime / date / time cells: parse_exam_time formats them itself.
    return value


def _raw_time(row: dict[str, Any]) -> str:
    parts = [str(row.get(name) or '') for name in ('time_text', 'date_text',
                                                   'start_text', 'end_text')]
    return ' '.join(part for part in parts if part)[:64]


def _limited(row: dict[str, Any], field: str) -> str:
    # A trimmed text cell cut to the model column length.
    return str(row.get(field) or '').strip()[:_TEXT_LIMITS[field]]


class Command(BaseCommand):
    help = ('Import the exam schedule table (xlsx or csv) into CourseExam rows '
            'of a term (see timetable/README.md §8.4).')

    def add_arguments(self, parser):
        parser.add_argument('file', help='path of the xlsx or csv file')
        parser.add_argument('--term', required=True,
                            help='AcademicTerm code (e.g. 26-27-1)')
        parser.add_argument('--sheet', default=None,
                            help='worksheet name (xlsx only; default: the first one)')
        parser.add_argument('--replace', action='store_true',
                            help="delete the term's existing exam rows first")

    def handle(self, *args, **options):
        term = AcademicTerm.objects.filter(code=options['term']).first()
        if term is None:
            raise CommandError(
                f'unknown term {options["term"]!r}; create it in the admin or '
                f'run timetable_seed_terms / import_academic_calendar first')
        path = Path(options['file'])
        if path.suffix.lower() == '.csv':
            if options['sheet']:
                raise CommandError('--sheet applies to xlsx files only')
            columns, rows = self._read_csv(path)
        else:
            columns, rows = self._read_xlsx(path, options['sheet'])
        self.stdout.write(f'{path.name}: {len(rows)} data row(s), '
                          f'columns {sorted(columns)}')

        prepared: dict[tuple[str, str, Any], dict[str, Any]] = {}
        skipped: list[str] = []
        other_terms: dict[str, int] = {}
        for number, row in rows:
            code = normalise_course_code(row.get('course_code'))
            name = str(row.get('name') or '').strip()
            if not code or not name:
                skipped.append(f'row {number}: missing 课程号/课程名')
                continue
            term_code = parse_term_code(row.get('term_text'))
            if term_code is not None and term_code != term.code:
                other_terms[term_code] = other_terms.get(term_code, 0) + 1
                continue
            parsed = parse_exam_time(
                row.get('time_text') or row.get('date_text') or '', term,
                start_text=row.get('start_text') or '',
                end_text=row.get('end_text') or '')
            if parsed is None:
                skipped.append(f'row {number}: unreadable time {_raw_time(row)!r}')
                continue
            start, end = parsed
            class_no = normalise_class_no(row.get('class_no'))
            key = (code[:32], class_no[:8], start)
            prepared[key] = {
                'name': name[:_TEXT_LIMITS['name']],
                'teacher': _limited(row, 'teacher'),
                'end': end,
                'room': _limited(row, 'room'),
                'method': _limited(row, 'method'),
                'note': _limited(row, 'note'),
                'raw_time': _raw_time(row),
            }

        created = updated = removed = 0
        with transaction.atomic():
            if options['replace']:
                removed, _ = CourseExam.objects.filter(term=term).delete()
            existing = {
                (exam.course_code, exam.class_no, exam.start): exam
                for exam in CourseExam.objects.select_for_update().filter(term=term)
            }
            for (code, class_no, start), fields in prepared.items():
                exam = existing.get((code, class_no, start))
                if exam is None:
                    CourseExam.objects.create(
                        term=term, course_code=code, class_no=class_no, start=start,
                        **fields)
                    created += 1
                    continue
                changed = [name for name, value in fields.items()
                           if getattr(exam, name) != value]
                if changed:
                    for name in changed:
                        setattr(exam, name, fields[name])
                    exam.save(update_fields=changed)
                    updated += 1
        for line in skipped:
            self.stdout.write(self.style.WARNING(f'skipped {line}'))
        summary = (f'done: created {created}, updated {updated}, removed {removed}; '
                   f'skipped {len(skipped)} row(s)')
        if other_terms:
            detail = ', '.join(f'{code} ({count})'
                               for code, count in sorted(other_terms.items()))
            summary += (f', {sum(other_terms.values())} row(s) of other term(s): '
                        f'{detail}')
        self.stdout.write(summary)

    def _read_xlsx(self, path: Path, sheet_name: str | None):
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (OSError, InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
            raise CommandError(f'cannot open workbook: {exc}')
        try:
            if sheet_name is None:
                if not workbook.sheetnames:
                    raise CommandError('the workbook has no worksheet')
                sheet = workbook[workbook.sheetnames[0]]
            elif sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
            else:
                raise CommandError(f'worksheet {sheet_name!r} not found; '
                                   f'available: {workbook.sheetnames}')
            return self._read_rows(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()

    def _read_csv(self, path: Path):
        try:
            with path.open(encoding='utf-8-sig', newline='') as handle:
                return self._read_rows(tuple(line) for line in csv.reader(handle))
        except OSError as exc:
            raise CommandError(f'cannot read {path}: {exc}')

    def _read_rows(
        self, lines,
    ) -> tuple[dict[str, int], list[tuple[int, dict[str, Any]]]]:
        columns: dict[str, int] | None = None
        rows: list[tuple[int, dict[str, Any]]] = []
        for number, values in enumerate(lines, start=1):
            if columns is None:
                columns = _column_map(tuple(values))
                if columns is None and number >= HEADER_SCAN_ROWS:
                    break
                continue
            if all(value is None or str(value).strip() == '' for value in values):
                continue
            row: dict[str, Any] = {}
            for field, index in columns.items():
                value = values[index] if index < len(values) else None
                row[field] = _cell_text(value)
            rows.append((number, row))
        if columns is None:
            raise CommandError(
                f'no header row with 课程号, 课程名 and 考试时间/考试日期 in the first '
                f'{HEADER_SCAN_ROWS} rows')
        return columns, rows
